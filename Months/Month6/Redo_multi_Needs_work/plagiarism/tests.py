import unittest
import os
import time
from fpdf import FPDF
from reportlab.lib.pagesizes import letter
from odf.opendocument import OpenDocumentText
from odf.text import P
from docx import Document
from ideal_completion import has_plagiarism

class TestPlagiarism(unittest.TestCase):
    
    @classmethod
    def setUpClass(cls):
        cls.plagiarism_text = """Donald Trump said he built the biggest, the broadest, the most unified coalition in all of American history. The Sky News data team assess whether he's right. This campaign, has been so historic in so many ways. We've built the biggest, the broadest, the most unified coalition."""
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(0, 0)
        pdf.set_font('arial', '', 13.0)
        # Split text into lines and write each line separately
        lines = cls.plagiarism_text.split('. ')
        for line in lines:
            pdf.cell(ln=1, h=5.0, align='L', w=0, txt=line+".", border=0)
        pdf.output('test.pdf', 'F')
        
        # write a docx file
        doc = Document()
        doc.add_paragraph(cls.plagiarism_text)
        doc.save("test.docx")
        
        # write odt
        textdoc = OpenDocumentText()
        para = P(text=cls.plagiarism_text)
        textdoc.text.addElement(para)
        textdoc.save("test.odt")
        
        cls.threshold = 0.9
        
    @classmethod
    def tearDownClass(cls):
        os.remove("test.pdf")
        os.remove("test.docx")
        os.remove("test.odt")
    
    def test_plagiarism_from_text(self):    
        time.sleep(5)
        self.assertTrue(has_plagiarism(self.plagiarism_text, self.threshold))
        
    def test_plagiarism_from_pdf(self):
        time.sleep(5)
        self.assertTrue(has_plagiarism("test.pdf", self.threshold))
    
    def test_plagiarism_from_docx(self):
        time.sleep(5)
        self.assertTrue(has_plagiarism("test.docx", self.threshold))
    
    def test_plagiarism_from_odt(self):
        time.sleep(5)
        self.assertTrue(has_plagiarism("test.odt", self.threshold))
    
    def test_no_plagiarism(self):
        time.sleep(5)
        doc = "John Doe is a singer"
        self.assertFalse(has_plagiarism(doc, self.threshold))
        
if __name__ == '__main__':
    unittest.main(verbosity=2)